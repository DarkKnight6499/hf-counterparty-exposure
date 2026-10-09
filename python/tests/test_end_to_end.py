"""End-to-end run, SA-CCR/CSA integration against hand values, README consistency and output hygiene."""
import csv
import math
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pytest import approx

from hfexposure import config as cfg
from hfexposure import pricers as px
from hfexposure.book import BookTrade, load_book, value_trade
from hfexposure.exposure import CsaRow, compute_exposure
from hfexposure.market import load_market
from hfexposure.memo import write_memos
from hfexposure.report_xlsx import write_workbook
from hfexposure.run import README_END, README_START, STRESS_END, STRESS_START, results_table, run_all, stress_table
from vendor.csa.terms import CSATerms
from vendor.saccr.trades import Trade


@pytest.fixture(scope="module")
def results():
    return run_all()


def test_book_shape_and_mtm_matches_fresh_repricing():
    m = load_market()
    book = load_book(m)                                  # raises if any CSV mtm differs from repricing
    counts = {}
    for bt in book:
        counts[bt.trade.netting_set] = counts.get(bt.trade.netting_set, 0) + 1
        assert abs(value_trade(bt, m) - bt.trade.mtm) <= cfg.PRICING_TOLERANCE_USD
    assert len(counts) == 3 and all(15 <= n <= 25 for n in counts.values())


def test_book_loader_rejects_a_tampered_mtm(tmp_path):
    rows = list(csv.DictReader(open(cfg.TRADES_CSV, newline="", encoding="utf-8")))
    rows[0][cfg.COL_MTM] = str(float(rows[0][cfg.COL_MTM]) + 5.0)
    bad = tmp_path / "trades.csv"
    with open(bad, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cfg.TRADE_COLUMNS)
        w.writeheader()
        w.writerows(rows)
    with pytest.raises(ValueError):
        load_book(load_market(), bad)


@pytest.mark.parametrize("mpor", [10, 20])
def test_saccr_ead_integration_hand_computed(mpor):
    # 5y par payer swap, USD 100m, no threshold/MTA/IA/collateral, margined with the given MPOR.
    # AddOn = 0.5% * N * SD * MF, SD = (1 - e^-0.25)/0.05, MF = 1.5 sqrt(MPOR/250); V = C = 0 gives multiplier 1;
    # EAD = 1.4 * AddOn.
    m = load_market()
    par = px.irs_par_rate(m.curves["USD"], 0, 5)
    t = Trade("T1", "NS-T", "IR", "USD", "Long", "Linear", None, 1e8, 0, 5, None, None, None, 0.0, "NA")
    bt = BookTrade(t, par, 0.0, 1e8)
    mtm = value_trade(bt, m)
    assert mtm == approx(0.0, abs=1e-6)
    csa = CsaRow(CSATerms("NS-T", "T", 0.0, 0.0, 0.0, 0.0, 0.0), "Daily", mpor, cfg.IM_REGIME_IA_ONLY)
    e = compute_exposure("NS-T", [t], mtm, csa)
    sd = (1 - math.exp(-0.25)) / 0.05
    addon = 0.005 * 1e8 * sd * 1.5 * math.sqrt(max(mpor, 10) / 250)
    assert e.saccr.addon.aggregate == approx(addon, rel=1e-12)
    assert e.saccr.ead == approx(1.4 * addon, rel=1e-9)
    assert e.saccr.rc == 0.0 and e.collateral == 0.0


def test_csa_call_and_collateral_bookkeeping(results):
    for r in results:
        m, t = r.metrics, r.csa.terms
        assert m.vm_held == approx(t.posted_vm + m.call.vm_call)
        assert m.im_held == approx(t.posted_im + m.call.im_call)
        assert m.saccr.c == approx(m.vm_held + m.im_held)
        if r.csa.im_regime == cfg.IM_REGIME_SIMM:
            assert m.im_held >= m.simm.total_im - r.csa.terms.mta
        else:
            assert m.im_held >= t.independent_amount


def test_end_to_end_invariants(results):
    assert [r.fund.netting_set for r in results] == ["NS-A", "NS-B", "NS-C"]
    for r in results:
        assert all(s.exposure >= 0 for s in r.stress)
        assert r.worst.exposure == max(s.exposure for s in r.stress)
        assert r.metrics.saccr.ead > 0 and r.metrics.simm.total_im > 0
        assert r.proposal.utilization == approx(r.worst.exposure / r.proposal.limit)
        assert len(r.stress) == len(cfg.SCENARIOS) + 2


def test_readme_results_table_matches_a_fresh_run(results):
    text = (cfg.ROOT_DIR / "README.md").read_text(encoding="utf-8")
    block = re.search(re.escape(README_START) + r"\n(.*?)\n" + re.escape(README_END), text, re.S).group(1)
    assert block == results_table(results)
    stress_block = re.search(re.escape(STRESS_START) + r"\n(.*?)\n" + re.escape(STRESS_END), text, re.S).group(1)
    assert stress_block == stress_table(results)


def test_workbook_and_memos(results, tmp_path):
    wb_path = tmp_path / "x.xlsx"
    write_workbook(results, wb_path)
    wb = load_workbook(wb_path)
    assert wb.properties.creator == "Yazad Madan"
    assert wb.sheetnames == ["Book", "Exposure", "Stress", "Limits"]
    for ws in wb:
        assert all(c.font.bold for c in ws[1])
        assert all(c.font.color is None for c in ws[1])
        for row in ws.iter_rows():
            for c in row:
                assert c.fill.fill_type is None
    assert wb["Book"].max_row - 1 == sum(len(r.trades) for r in results)
    paths = write_memos(results, tmp_path)
    assert len(paths) == 3
    for p in paths:
        text = p.read_text(encoding="utf-8")
        for needle in ("Limit proposal", "Stress tests", "Limitations", "Recommendation"):
            assert needle in text
        assert len(text.splitlines()) < 60            # one page


def test_no_em_dashes_or_attribution_in_repo_text():
    root = cfg.ROOT_DIR
    banned = ["—", "–", "Co-Authored-By", "Claude"]
    for path in root.rglob("*"):
        if not path.is_file() or ".git" in path.parts or "__pycache__" in path.parts or ".pytest_cache" in path.parts:
            continue
        if path.suffix not in {".py", ".md", ".csv", ".json", ".ini", ".txt"}:
            continue
        if path.name == "test_end_to_end.py":
            continue
        text = path.read_text(encoding="utf-8")
        for b in banned:
            assert b not in text, f"{b!r} in {path}"
